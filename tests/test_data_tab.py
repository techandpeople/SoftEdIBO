"""The Data tab: session list, views, corrections and the Excel export."""

import json
from datetime import datetime, timedelta

import pytest

from src.data import interaction_model as im
from src.data.database import Database
from src.data.models import (
    EventCorrection,
    InteractionEvent,
    ParticipantRecord,
    SessionAnnotation,
    SessionAssignment,
    SessionRecord,
)
from src.data.research_export import ExportOptions, ResearchExporter
from src.data.session_data import load_session_data
from src.data.session_index import NO_GROUP, build_session_rows
from src.data.timeline import KIND_PHASE, KIND_RESPONSE, KIND_ROBOT, build_timeline
from src.gui.data_panel import DataPanel

START = datetime(2026, 10, 12, 10, 0, 0)


@pytest.fixture
def db(tmp_path):
    database = Database(str(tmp_path / "test.db"))
    database.connect()
    yield database
    database.close()


def _event(db, session_id, t_ms, type_, action, target="", participant="system",
           **meta):
    db.log_event(InteractionEvent(
        session_id=session_id, participant_id=participant, type=type_,
        action=action, target=target,
        timestamp=START + timedelta(milliseconds=t_ms),
        metadata=json.dumps(meta) if meta else ""))


def _seed(db, session_id="S001", group="", status="valid", finished=True):
    """A two-child, two-robot session: P1 leads P2 by 200 ms, five times."""
    db.save_session(SessionRecord(session_id, "Any activity", START,
                                  START + timedelta(seconds=30)))
    for pid, robot, skin in (("P1", "thymio-1", "skin-1"),
                             ("P2", "thymio-2", "skin-2")):
        db.save_participant(ParticipantRecord(pid, f"child {pid}", 8))
        db.link_participant_to_session(session_id, pid)
        db.save_assignment(SessionAssignment(session_id, robot, pid, [skin]))
    _event(db, session_id, 0, "session", "start",
           robot_kinds={"thymio-1": "thymio", "thymio-2": "thymio"})
    _event(db, session_id, 0, "activity", "state", "thymio-1/skin-1",
           **{"from": "", "to": "phase1", "final": False,
              "trigger": {"source": "start"}})
    _event(db, session_id, 500, "robot", "inflate", "thymio-1/skin-1",
           cause="phase", params={})
    for t in range(1000, 6000, 1000):
        for offset, pid, skin in ((0, "P1", "skin-1"), (200, "P2", "skin-2")):
            _event(db, session_id, t + offset, "touch", "press", f"{skin}:0",
                   pid, sensor_id=0, intensity_ut=150.0)
            _event(db, session_id, t + offset + 120, "touch", "release",
                   f"{skin}:0", pid, sensor_id=0, duration_ms=120,
                   peak_intensity_ut=300.0)
    _event(db, session_id, 1050, "robot", "set_led", "thymio-1/skin-1",
           cause="touch", params={})
    if finished:
        _event(db, session_id, 7000, "activity", "state", "thymio-1/skin-1",
               **{"from": "phase1", "to": "done", "final": True,
                  "in_state_ms": 7000,
                  "trigger": {"source": "condition",
                              "condition": {"touch_count": {"min": 5}}}})
    db.flush_events()
    if group or status != "valid":
        db.save_session_annotation(SessionAnnotation(
            session_id, status=status, group_label=group))


# --- session index -----------------------------------------------------------

def test_session_rows_summarise_each_session(db):
    _seed(db, "S001", group="G1")
    _seed(db, "S002", finished=False)
    rows = {r.session_id: r for r in build_session_rows(db)}
    assert (rows["S001"].group, rows["S001"].finished, rows["S001"].children) == (
        "G1", "1/1", 2)
    assert rows["S002"].finished == "0/1"
    assert rows["S001"].duration_s == 30.0


def test_session_row_filters_by_text_and_status(db):
    _seed(db, "S001", group="G1", status="test")
    [row] = build_session_rows(db)
    assert row.matches("g1", "") and row.matches("child p2", "")
    assert row.matches("", "test") and not row.matches("", "valid")
    assert not row.matches("turtle", "")


# --- timeline ----------------------------------------------------------------

def test_timeline_has_phases_a_lane_per_child_and_robot_marks(db):
    _seed(db)
    data = load_session_data(db, "S001")
    assert data is not None
    timeline = build_timeline(data, im.InteractionModel(data))
    lanes = {lane.label: lane for lane in timeline.lanes}
    assert list(lanes) == ["Phases: thymio-1/skin-1", "P1", "P2", "Robot",
                           "Session"]
    phases = lanes["Phases: thymio-1/skin-1"].items
    assert [(p.label, p.start_ms, p.end_ms, p.kind) for p in phases] == [
        ("phase1", 0.0, 7000.0, KIND_PHASE), ("done", 7000.0, 30000.0, KIND_PHASE)]
    first = lanes["P1"].items[0]
    assert (first.start_ms, first.end_ms) == (1000.0, 1120.0)
    assert "peak 300 uT" in first.label
    assert [item.kind for item in lanes["Robot"].items] == [
        KIND_ROBOT, KIND_RESPONSE]
    assert timeline.duration_ms == 30000.0


# --- export ------------------------------------------------------------------

def test_export_builds_every_sheet_with_a_dictionary(db):
    _seed(db, "S001", group="G1")
    tables = ResearchExporter(db).build_tables(["S001"])
    names = [t.name for t in tables]
    assert names[:7] == ["Sessions", "Milestones", "Participants", "Touches",
                         "Profiles", "Pairs", "Events"]
    assert names[7:11] == ["Episodes", "Turns", "Behaviours", "Reactions"]
    assert names[-1] == "Dictionary"
    by_name = {t.name: t for t in tables}
    assert len(by_name["Touches"].rows) == 10
    assert [(r["participant_id"], r["group"], r["robot_id"])
            for r in by_name["Participants"].rows] == [
        ("P1", "G1", "thymio-1"), ("P2", "G1", "thymio-2")]
    documented = {(r["sheet"], r["column"]) for r in by_name["Dictionary"].rows}
    for table in tables[:-1]:
        for column in table.columns:
            assert (table.name, column.name) in documented


def test_export_skips_test_and_excluded_sessions_unless_asked(db):
    _seed(db, "S001")
    _seed(db, "S002", status="test")
    _seed(db, "S003", status="excluded")
    exporter = ResearchExporter(db)
    ids = ["S001", "S002", "S003"]

    def exported(options):
        sessions = exporter.build_tables(ids, options)[0]
        return [row["session_id"] for row in sessions.rows]

    assert exported(ExportOptions()) == ["S001"]
    assert exported(ExportOptions(include_test=True)) == ["S001", "S002"]
    assert exported(ExportOptions(include_test=True, include_excluded=True)) == ids


def test_export_keeps_raw_events_and_marks_corrections(db):
    _seed(db)
    press = next(e for e in db.get_session_events("S001")
                 if e.type == "touch" and e.action == "press")
    db.save_event_correction(EventCorrection(
        press.event_id, "S001", excluded=True, note="adult"))
    tables = {t.name: t for t in ResearchExporter(db).build_tables(["S001"])}
    assert len(tables["Touches"].rows) == 9          # left out of the measures
    marked = [r for r in tables["Events"].rows if r["excluded"]]
    assert [(r["event_id"], r["correction_note"]) for r in marked] == [
        (press.event_id, "adult")]


def test_export_writes_a_readable_workbook(db, tmp_path):
    from openpyxl import load_workbook
    _seed(db, "S001", group="G1")
    path = tmp_path / "out.xlsx"
    assert ResearchExporter(db).export(["S001"], path) == 1
    workbook = load_workbook(path)
    assert workbook.sheetnames[0] == "Sessions" and "Dictionary" in workbook.sheetnames
    touches = workbook["Touches"]
    headers = [cell.value for cell in touches[1]]
    first = dict(zip(headers, [cell.value for cell in touches[2]]))
    assert (first["source"], first["t_ms"], first["duration_ms"]) == ("P1", 1000, 120)
    assert first["time"] == "2026-10-12T10:00:01.000"
    assert touches.freeze_panes == "A2"


def test_export_can_group_by_skin(db):
    _seed(db)
    tables = {t.name: t for t in ResearchExporter(db).build_tables(
        ["S001"], ExportOptions(source_key=im.SOURCE_SKIN))}
    assert sorted(r["source"] for r in tables["Profiles"].rows) == [
        "skin-1", "skin-2"]


# --- panel -------------------------------------------------------------------

def _cell(table, row, column) -> str:
    item = table.item(row, column)
    return item.text() if item is not None else ""


def _panel(qtbot, db):
    panel = DataPanel(db)
    qtbot.addWidget(panel)
    return panel


def test_panel_lists_sessions_under_their_group(qtbot, db):
    _seed(db, "S001", group="G1")
    _seed(db, "S002")
    panel = _panel(qtbot, db)
    assert set(panel.group_labels()) == {"G1", NO_GROUP}
    assert panel.listed_session_ids("G1") == ["S001"]
    assert panel.session_cell("S001", 6) == "1/1"


def test_panel_lists_newest_session_first(qtbot, db):
    db.save_session(SessionRecord("s01", "A", datetime(2026, 1, 1, 9, 0)))
    db.save_session(SessionRecord("s02", "A", datetime(2026, 1, 1, 10, 0),
                                  datetime(2026, 1, 1, 10, 30)))
    panel = _panel(qtbot, db)
    assert panel.listed_session_ids() == ["s02", "s01"]
    assert panel.session_cell("s02", 3) == "30:00"


def test_panel_filter_narrows_the_list(qtbot, db):
    _seed(db, "S001", group="G1")
    _seed(db, "S002", status="test")
    panel = _panel(qtbot, db)
    panel.filter_edit.setText("g1")
    assert panel.listed_session_ids() == ["S001"]
    panel.filter_edit.setText("")
    panel.status_filter.setCurrentIndex(panel.status_filter.findData("test"))
    assert panel.listed_session_ids() == ["S002"]


def test_selecting_a_session_fills_every_view(qtbot, db):
    _seed(db)
    panel = _panel(qtbot, db)
    assert panel.select_session("S001")
    assert panel.summary_view.activity_value.text() == "Any activity"
    assert "own robot each" in panel.summary_view.robots_value.text()
    assert "1 of 1 finished" in panel.summary_view.outcome_value.text()
    assert panel.children_view.profiles_table.rowCount() == 2
    assert panel.children_view.pairs_table.rowCount() == 1
    assert panel.events_view.events_table.rowCount() == len(
        db.get_session_events("S001"))
    low, high = panel.timeline_view.canvas.visible_range_ms()
    assert (low, high) == (0.0, 30000.0)


def test_source_combo_regroups_the_children_view(qtbot, db):
    _seed(db)
    panel = _panel(qtbot, db)
    panel.select_session("S001")
    panel.source_combo.setCurrentIndex(panel.source_combo.findData(im.SOURCE_SKIN))
    table = panel.children_view.profiles_table
    assert sorted(_cell(table, r, 0) for r in range(table.rowCount())) == [
        "skin-1", "skin-2"]


def test_saving_the_review_stores_it_and_regroups_the_list(qtbot, db):
    _seed(db)
    panel = _panel(qtbot, db)
    panel.select_session("S001")
    summary = panel.summary_view
    summary.group_edit.setText("G7")
    summary.order_spin.setValue(2)
    summary.status_combo.setCurrentIndex(summary.status_combo.findData("test"))
    summary.notes_edit.setPlainText("pilot run")
    summary.save_btn.click()

    saved = db.get_session_annotation("S001")
    assert (saved.group_label, saved.visit_order, saved.status, saved.notes) == (
        "G7", 2, "test", "pilot run")
    assert panel.group_labels() == ["G7"]
    assert panel.selected_session_ids() == ["S001"]


def test_excluding_and_restoring_events_from_the_events_view(qtbot, db):
    _seed(db)
    panel = _panel(qtbot, db)
    panel.select_session("S001")
    view = panel.events_view
    view.type_filter.setCurrentIndex(view.type_filter.findData("touch"))
    assert view.events_table.rowCount() == 20
    view.events_table.selectRow(0)
    [event_id] = view.selected_event_ids()

    view.exclude_btn.click()
    assert db.get_event_corrections("S001")[event_id].excluded
    assert panel.children_view.profiles_table.rowCount() == 2
    assert _cell(view.events_table, 0, 7) == "excluded"

    view.events_table.selectRow(0)
    view.clear_btn.click()
    assert db.get_event_corrections("S001") == {}


def test_reattributing_an_event_to_another_child(qtbot, db):
    _seed(db)
    panel = _panel(qtbot, db)
    panel.select_session("S001")
    view = panel.events_view
    view.type_filter.setCurrentIndex(view.type_filter.findData("touch"))
    view.events_table.selectRow(0)
    [event_id] = view.selected_event_ids()
    view.participant_combo.setCurrentIndex(view.participant_combo.findData("P2"))
    view.reassign_btn.click()
    assert db.get_event_corrections("S001")[event_id].participant_id == "P2"
    assert "attributed to P2" in _cell(view.events_table, 0, 7)


def test_delete_moves_selected_session_to_trash(qtbot, db, monkeypatch):
    from PySide6.QtWidgets import QMessageBox
    _seed(db)
    panel = _panel(qtbot, db)
    assert not panel.delete_btn.isEnabled()
    panel.select_session("S001")
    assert panel.delete_btn.isEnabled()
    monkeypatch.setattr(QMessageBox, "exec",
                        lambda self: QMessageBox.StandardButton.Yes)
    panel._on_delete()
    assert panel.listed_session_ids() == []
    assert db.get_all_sessions() == []
    assert [t.session_id for t in db.list_trashed_sessions()] == ["S001"]
    assert panel.summary_view.activity_value.text() == "-"


def test_export_dialog_counts_and_writes(qtbot, db, tmp_path):
    from src.gui.research_export_dialog import ResearchExportDialog
    _seed(db, "S001")
    _seed(db, "S002", status="test")
    rows = build_session_rows(db)
    dialog = ResearchExportDialog(db, rows, ["S002"])
    qtbot.addWidget(dialog)
    assert dialog.selected_radio.isChecked()
    assert dialog.session_ids() == [] and not dialog.export_btn.isEnabled()
    dialog.include_test_check.setChecked(True)
    assert dialog.session_ids() == ["S002"]
    dialog.all_radio.setChecked(True)
    assert dialog.session_ids() == ["S001", "S002"]
    assert dialog.export_to(tmp_path / "all.xlsx") == 2


def test_timeline_canvas_zooms_and_describes_what_is_under_the_cursor(qtbot, db):
    from PySide6.QtCore import QPoint
    _seed(db)
    panel = _panel(qtbot, db)
    panel.resize(1300, 700)
    panel.select_session("S001")
    panel.views_tabs.setCurrentWidget(panel.timeline_view)
    panel.show()
    qtbot.waitExposed(panel)
    canvas = panel.timeline_view.canvas
    canvas.zoom(2.0)
    low, high = canvas.visible_range_ms()
    assert high - low == 15000.0
    panel.timeline_view.fit_btn.click()
    assert canvas.visible_range_ms() == (0.0, 30000.0)
    # Lane 0 is the phase lane: its first phase covers the left of the plot.
    assert "phase1" in canvas.describe_at(QPoint(200, 22 + 13))
    assert canvas.describe_at(QPoint(200, 5)) == ""


def test_patterns_view_pools_the_selected_sessions(qtbot, db):
    _seed(db, "S001", group="G1")
    _seed(db, "S002", group="G1")
    panel = _panel(qtbot, db)
    panel.select_session("S001")
    view = panel.patterns_view
    assert view.turns_table.rowCount() == 0          # not computed while hidden
    panel.views_tabs.setCurrentWidget(view)
    assert "1 selected" in view.scope_label.text()
    assert view.turns_table.rowCount() >= 1
    assert view.reactions_table.rowCount() == 1

    panel.sessions_tree.selectAll()                  # both sessions and the group
    assert "2 selected" in view.scope_label.text()
    assert _cell(view.reactions_table, 0, 6) == "2"  # rests on two sessions
