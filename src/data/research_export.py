"""Excel export of sessions for statistical analysis.

Writes one workbook with a sheet per table (sessions, milestones,
participants, touches, profiles, pairs, events) plus a dictionary sheet that
explains every column, so the file can be handed to someone who never saw
the app.

Building the tables is separate from writing the file: ``build_tables`` is
pure data and is what the tests check; ``export`` needs ``openpyxl``.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Iterable

from src.data import interaction_model as im
from src.data.export import EventAttribution
from src.data.interaction_dynamics import InteractionDynamics
from src.data.interaction_model import Column, Table
from src.data.models import SESSION_STATUS_EXCLUDED, SESSION_STATUS_TEST
from src.data.session_data import SessionData, load_session_data
from src.ml import interaction_clusters as clusters

if TYPE_CHECKING:
    from src.data.database import Database

PARTICIPANT_COLUMNS = (
    Column("session_id", "Session."),
    Column("group", "Group of children (key to join with the questionnaires)."),
    Column("participant_id", "Participant code (key to join with the "
           "questionnaires)."),
    Column("alias", "Name or nickname recorded for the participant."),
    Column("age", "Age.", "years"),
    Column("robot_id", "Robot the participant was assigned to."),
    Column("units", "Skins or branches assigned to the participant."),
)

EVENT_COLUMNS = (
    Column("session_id", "Session."),
    Column("event_id", "Identifier of the event in the database."),
    Column("time", "Wall-clock time of the event.", "ISO 8601"),
    Column("t_ms", "Time since the session started.", "ms"),
    Column("participant_id", "Participant as recorded."),
    Column("robot_id", "Robot the event belongs to."),
    Column("type", "Kind of event: touch, activity, robot, input, gesture, "
           "sync, observer, marker, session."),
    Column("action", "What happened (press, release, state, a robot action, "
           "an input name, an observer code, ...)."),
    Column("target", "What it happened on (skin:chamber or robot/skin)."),
    Column("details", "Everything else recorded with the event.", "JSON"),
    Column("corrected_participant", "Participant the researcher attributed "
           "the event to, when corrected."),
    Column("excluded", "True when the researcher left the event out of the "
           "analysis."),
    Column("correction_note", "The researcher's note on the correction."),
)

DICTIONARY_COLUMNS = (
    Column("sheet", "Sheet the column is on."),
    Column("column", "Column name."),
    Column("meaning", "What the column holds."),
    Column("unit", "Unit or format of the values."),
)


@dataclass(frozen=True)
class ExportOptions:
    """What to include and how to group touches."""

    source_key: str = im.SOURCE_PARTICIPANT
    include_test: bool = False
    include_excluded: bool = False

    def accepts(self, status: str) -> bool:
        """Whether a session with this review status is exported."""
        if status == SESSION_STATUS_TEST:
            return self.include_test
        if status == SESSION_STATUS_EXCLUDED:
            return self.include_excluded
        return True


def _participants_table(data: SessionData) -> Table:
    group = data.annotation.group_label if data.annotation else ""
    assigned: dict[str, list] = {}
    for assignment in data.assignments:
        assigned.setdefault(assignment.participant_id, []).append(assignment)
    rows = []
    for participant in data.participants:
        for assignment in assigned.get(participant.participant_id) or [None]:
            rows.append({
                "session_id": data.session.session_id,
                "group": group,
                "participant_id": participant.participant_id,
                "alias": participant.alias,
                "age": participant.age,
                "robot_id": assignment.robot_id if assignment else "",
                "units": ", ".join(assignment.unit_ids) if assignment else "",
            })
    return Table("Participants", PARTICIPANT_COLUMNS, rows)


def _events_table(data: SessionData) -> Table:
    attribution = EventAttribution(data.assignments)
    start = data.session.start_time
    rows = []
    for event in sorted(data.events, key=lambda e: e.timestamp):
        correction = (data.corrections.get(event.event_id)
                      if event.event_id is not None else None)
        robot_id, _participant = attribution.resolve(event)
        rows.append({
            "session_id": data.session.session_id,
            "event_id": event.event_id,
            "time": event.timestamp.isoformat(timespec="milliseconds"),
            "t_ms": round((event.timestamp - start).total_seconds() * 1000.0),
            "participant_id": event.participant_id,
            "robot_id": robot_id,
            "type": event.type,
            "action": event.action,
            "target": event.target,
            "details": event.metadata,
            "corrected_participant":
                (correction.participant_id or "") if correction else "",
            "excluded": bool(correction and correction.excluded),
            "correction_note": correction.note if correction else "",
        })
    return Table("Events", EVENT_COLUMNS, rows)


def session_tables(data: SessionData, options: ExportOptions) -> list[Table]:
    """Every exported table of one session."""
    model = im.InteractionModel(data, options.source_key)
    return [model.session(), model.milestones(), _participants_table(data),
            model.touches(), model.profiles(), model.pairs(),
            _events_table(data)]


def pooled_tables(sessions: list[SessionData],
                  options: ExportOptions) -> list[Table]:
    """The tables pooled over all exported sessions: touching episodes,
    turn-taking, behaviour transitions, reactions and, when the optional ML
    packages are installed, the touching styles (written on each episode)."""
    dynamics = InteractionDynamics(sessions, options.source_key)
    episodes = dynamics.episodes()
    tables = [episodes, dynamics.turns(), dynamics.behaviours(),
              dynamics.reactions()]
    try:
        result = clusters.cluster_episodes(episodes)
    except clusters.MLNotInstalled:
        return tables
    for row, label in zip(episodes.rows, result.labels):
        row["style"] = label
    return tables + [result.summary]


def dictionary_table(tables: list[Table]) -> Table:
    """The sheet that explains every column of ``tables``."""
    rows = [{"sheet": table.name, "column": column.name,
             "meaning": column.description, "unit": column.unit}
            for table in tables for column in table.columns]
    return Table("Dictionary", DICTIONARY_COLUMNS, rows)


class ResearchExporter:
    """Exports sessions to an Excel workbook.

    Args:
        db: Open database to read the sessions from.
    """

    def __init__(self, db: "Database") -> None:
        self._db = db

    def build_tables(self, session_ids: Iterable[str],
                     options: ExportOptions | None = None) -> list[Table]:
        """The workbook's tables for ``session_ids`` (sessions whose review
        status the options leave out are skipped), dictionary last."""
        options = options or ExportOptions()
        sessions = []
        for session_id in session_ids:
            data = load_session_data(self._db, session_id)
            if data is None:
                continue
            status = data.annotation.status if data.annotation else ""
            if options.accepts(status):
                sessions.append(data)
        if not sessions:
            return []
        tables = im.combine([session_tables(data, options) for data in sessions])
        tables += pooled_tables(sessions, options)
        return tables + [dictionary_table(tables)]

    def export(self, session_ids: Iterable[str], path: str | Path,
               options: ExportOptions | None = None) -> int:
        """Write the workbook to ``path``. Returns the sessions exported."""
        tables = self.build_tables(session_ids, options)
        write_workbook(tables, path)
        sessions = next((t for t in tables if t.name == "Sessions"), None)
        return len(sessions.rows) if sessions else 0


def _cell(value):
    """A value Excel can hold (dicts/lists become JSON text)."""
    if isinstance(value, (dict, list)):
        return json.dumps(value, sort_keys=True)
    return value


def write_workbook(tables: list[Table], path: str | Path) -> None:
    """Write ``tables`` as the sheets of an ``.xlsx`` workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    for default_sheet in list(workbook.worksheets):
        workbook.remove(default_sheet)
    for table in tables:
        sheet = workbook.create_sheet(table.name)
        sheet.append(table.headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for values in table.values():
            sheet.append([_cell(value) for value in values])
        sheet.freeze_panes = "A2"
        if table.rows:
            sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(table.columns, start=1):
            longest = max(
                [len(column.name)] + [len(str(row.get(column.name) or ""))
                                      for row in table.rows[:200]])
            sheet.column_dimensions[get_column_letter(index)].width = min(
                60, longest + 2)
    if not tables:
        workbook.create_sheet("Sessions")
    workbook.save(str(path))
